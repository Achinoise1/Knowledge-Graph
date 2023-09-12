import re
from docx import Document

path = '../教学大纲.docx'

try:    
    doc = Document(path)
except Exception as e:
    print("Error!",e)
else:
    print("Success!")      
    tables = doc.tables  
    table = tables[1]
    
count = -1

DicOrigin = {}

for j in range(1,len(table.rows)):                                       
    Order = table.cell(j,0).text                            
    Content = table.cell(j,1).text.split("\n")             
    DicOrigin.update({Order:Content})      
print(DicOrigin)


'''
Test Over 
即便一个表格分散在几个页面中
python一样认为这是一个表格
因此！
只要找到表头
往下第一个表格就是目标表格
修改之后
时间复杂度仍然不变
'''